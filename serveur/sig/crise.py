"""Carte de gestion de crise (lot C1) : éléments touchés autour d'un événement, zone par zone, en Lambert 93.

Principe : pour un point d'événement, chaque objet des couches analysées (section « crise » de config/sig.yaml)
reçoit sa distance exacte au point (point le plus proche pour une ligne ou une surface), jusqu'à « rayon_max_m ».
Un objet est dans un cercle de rayon r si cette distance est ≤ r : l'interface recalcule ainsi les compteurs en
direct pendant que l'utilisateur tire le cercle, sans appel au serveur. Une zone dessinée (polygone) est croisée
ici. Un élément appartient à la zone la plus grave qui le contient (ordre de « crise.zones »).

Rôles des catégories :
- « points » (défaut) : un élément par objet, ou par valeur du champ « regroupement » (une station = ses lignes) ;
- « stations » : comme « points », et les accès de la catégorie « acces » sont rattachés à la station la plus
  proche (« rattachement_acces_m ») : une station est touchée par son point ou par l'un de ses accès ;
- « lignes » : objets regroupés (une ligne = ses tronçons), longueur dans la zone, partie touchée dessinée ;
- « plai » : points des PLAI ; un point de rassemblement dans une zone bloquante est déclaré inutilisable et le
  plus proche hors de ces zones est proposé ;
- « maquettes » : maquettes positionnées regroupées par ouvrage (dossier), ouvrages annexes signalés comme puits.

Ne jamais bloquer : une couche absente ou illisible devient une anomalie, les autres catégories sont traitées.
"""
from __future__ import annotations

import getpass
import io
import json
import math
import re
import threading
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

import numpy as np
import shapely
from pyproj import Transformer
from shapely.geometry import LineString, Point, Polygon, mapping
from shapely.ops import nearest_points, unary_union

from . import crise_externe, crise_local
from . import inventaire as inv_sig
from . import objets as obj_sig

_VERS_L93 = Transformer.from_crs(4326, 2154, always_xy=True)
_VERS_WGS84 = Transformer.from_crs(2154, 4326, always_xy=True)
ROLES = ("points", "stations", "acces", "lignes", "plai", "maquettes", "secours", "rues", "bouclage", "denombrement")
ROLES_LOCAUX = ("rues", "bouclage", "denombrement")   # objets d'une source autour de l'événement (crise_local)
TAMPON_MAX_M = 2000       # couloir : tampon de part et d'autre de l'axe
FUSION_BOUCLAGE_M = 30     # points de bouclage d'une même voie plus proches que cela : un seul point
PAS_COURBE_M = 25          # pas des courbes de dénombrement (compteurs en direct pendant le glissement)
DIRECTIONS = ("N", "NE", "E", "SE", "S", "SO", "O", "NO")
CACHE_CANDIDATS = 16       # points d'événement gardés en mémoire (technique)
SEGMENTS_CERCLE = 64       # quarts de cercle des zones circulaires (écart au vrai cercle < 8 cm à 1 km)
DECIMALES_WGS84 = 6        # ~10 cm
STATUTS_DEFAUT = [{"id": "propose", "libelle": "À décider", "decide": False},
                  {"id": "ferme", "libelle": "Fermé", "decide": True, "barre": True},
                  {"id": "surveille", "libelle": "À surveiller", "decide": True},
                  {"id": "leve", "libelle": "Levé", "decide": True}]
MAX_JOURNAL = 5000         # entrées de main courante gardées (technique)

_verrou = threading.Lock()
_objets_couches: dict = {}                    # (id couche, version) -> objets en Lambert 93
_candidats: OrderedDict = OrderedDict()  # (centre, rayon, versions) -> candidats
_parts_courbes: dict = {}                # (catégorie, centre, objets) -> parts des surfaces par rayon (dénombrement)


class ErreurCrise(ValueError):
    """Demande invalide (zone mal formée, point hors du domaine) ; le message dit quoi corriger."""


# ---------------------------------------------------------------- configuration

def controler(c, erreur) -> None:
    """Section « crise » de config/sig.yaml (facultative)."""
    def nombre(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0

    if not isinstance(c, dict):
        raise erreur("« crise » doit être une section (dictionnaire)")
    for cle in ("rayon_max_m", "rayon_defaut_m", "pas_m", "rattachement_acces_m"):
        if not nombre(c.get(cle)):
            raise erreur(f"« crise.{cle} » doit être un nombre strictement positif")
    if c["rayon_defaut_m"] > c["rayon_max_m"]:
        raise erreur("« crise.rayon_defaut_m » doit être inférieur ou égal à « rayon_max_m »")
    if not isinstance(c.get("dossier_travail"), str) or not c["dossier_travail"]:
        raise erreur("« crise.dossier_travail » attendu (dossier des événements enregistrés, hors Git)")
    r = c.get("raccourcis_m")
    if not (isinstance(r, list) and r and all(nombre(v) and v <= c["rayon_max_m"] for v in r)):
        raise erreur("« crise.raccourcis_m » doit être une liste de rayons (m) entre 0 et « rayon_max_m »")
    if not (isinstance(c.get("types"), list) and c["types"] and all(isinstance(t, str) for t in c["types"])):
        raise erreur("« crise.types » doit être une liste de types d'événement")
    zones = c.get("zones")
    if not (isinstance(zones, list) and zones):
        raise erreur("« crise.zones » doit être une liste de rôles de zone (du plus grave au moins grave)")
    ids = set()
    for i, z in enumerate(zones, 1):
        if not (isinstance(z, dict) and isinstance(z.get("id"), str) and re.fullmatch(r"[a-z0-9_]+", z["id"])
                and isinstance(z.get("libelle"), str) and isinstance(z.get("couleur"), str)):
            raise erreur(f"« crise.zones » n° {i} : « id » (minuscules, chiffres, _), « libelle » et « couleur » attendus")
        if z["id"] in ids:
            raise erreur(f"« crise.zones » : identifiant « {z['id']} » en double")
        ids.add(z["id"])
    crise_local.controler(c, erreur)
    crise_externe.controler(c, erreur)
    ids_locales = {s["id"] for s in c.get("sources_locales", [])}
    cats = c.get("categories")
    if not (isinstance(cats, list) and cats):
        raise erreur("« crise.categories » doit être une liste de catégories analysées")
    vus = set()
    for i, k in enumerate(cats, 1):
        nom = f"« crise.categories » n° {i}"
        if not (isinstance(k, dict) and isinstance(k.get("id"), str) and isinstance(k.get("libelle"), str)):
            raise erreur(f"{nom} : « id » et « libelle » attendus")
        if k["id"] in vus:
            raise erreur(f"{nom} : identifiant « {k['id']} » en double")
        vus.add(k["id"])
        role = k.get("role", "points")
        if role not in ROLES:
            raise erreur(f"{nom} : « role » doit être l'un de : {', '.join(ROLES)}")
        if "source_locale" in k:
            if k["source_locale"] not in ids_locales:
                raise erreur(f"{nom} : « source_locale » doit citer une source de « crise.sources_locales » "
                             f"({', '.join(sorted(ids_locales)) or 'aucune déclarée'})")
            if "couches" in k:
                raise erreur(f"{nom} : « couches » ou « source_locale », pas les deux")
        elif role in ROLES_LOCAUX:
            raise erreur(f"{nom} : le rôle « {role} » demande « source_locale » (source autour de l'événement)")
        elif role != "maquettes" and not (isinstance(k.get("couches"), list) and k["couches"]
                                          and all(isinstance(x, str) for x in k["couches"])):
            raise erreur(f"{nom} : « couches » attendu (libellés exacts des couches, comme dans le gestionnaire)")
        if role == "bouclage":
            zb = k.get("zones_bouclage")
            if not (isinstance(zb, list) and zb and all(z in ids for z in zb)):
                raise erreur(f"{nom} : « zones_bouclage » attendu (rôles de zone dont la limite est bouclée)")
        if role == "denombrement":
            m = k.get("mesures")
            if not (isinstance(m, list) and m and all(isinstance(x, dict) and isinstance(x.get("id"), str)
                                                       and isinstance(x.get("libelle"), str) for x in m)):
                raise erreur(f"{nom} : « mesures » attendu (id, libelle ; champ facultatif : nombre d'objets sinon)")
        p = k.get("propositions", {})
        if not isinstance(p, dict) or any(z not in ids for z in p):
            raise erreur(f"{nom} : « propositions » doit associer un texte à des rôles de zone déclarés ({', '.join(ids)})")
        if role == "plai":
            b = k.get("zones_bloquantes", [])
            if not isinstance(k.get("champ_type"), str) or not isinstance(k.get("type_rassemblement"), str) \
                    or not isinstance(b, list) or any(z not in ids for z in b):
                raise erreur(f"{nom} : « champ_type », « type_rassemblement » et « zones_bloquantes » (rôles de zone) attendus")
        if role == "maquettes":
            d = k.get("dossiers")
            if not (isinstance(d, dict) and d and all(isinstance(v, str) for v in d.values())):
                raise erreur(f"{nom} : « dossiers » attendu (nom du dossier de la maquette → type d'ouvrage)")
        if role == "secours":
            if not (isinstance(k.get("nombre", 1), int) and not isinstance(k.get("nombre", 1), bool) and 1 <= k.get("nombre", 1) <= 10):
                raise erreur(f"{nom} : « nombre » doit être un entier de 1 à 10 (sites les plus proches montrés)")
            if "rayon_recherche_m" in k and not (nombre(k["rayon_recherche_m"]) and k["rayon_recherche_m"] <= 50000):
                raise erreur(f"{nom} : « rayon_recherche_m » doit être compris entre 0 et 50 000 m")
        if "couleur" in k and not isinstance(k["couleur"], str):
            raise erreur(f"{nom} : « couleur » doit être une couleur (ex. « #1c3f94 »)")
    roles = [k.get("role", "points") for k in cats]
    if "acces" in roles and "stations" not in roles:
        raise erreur("« crise.categories » : une catégorie « acces » demande une catégorie « stations »")
    statuts = c.get("statuts", STATUTS_DEFAUT)
    if not (isinstance(statuts, list) and statuts and all(isinstance(x, dict) and isinstance(x.get("id"), str)
                                                          and isinstance(x.get("libelle"), str) for x in statuts)):
        raise erreur("« crise.statuts » doit être une liste de statuts (id, libelle ; decide, barre facultatifs)")
    ids_statuts = [x["id"] for x in statuts]
    if len(set(ids_statuts)) != len(ids_statuts):
        raise erreur("« crise.statuts » : identifiant en double")
    if statuts[0].get("decide"):
        raise erreur("« crise.statuts » : le premier statut est celui « à décider » (decide: false)")
    for k in cats:
        if "statuts" in k and not (isinstance(k["statuts"], list) and k["statuts"] and k["statuts"][0] == ids_statuts[0]
                                   and all(x in ids_statuts for x in k["statuts"])):
            raise erreur(f"« crise.categories » {k['id']} : « statuts » doit commencer par « {ids_statuts[0]} » "
                         f"et ne citer que des statuts déclarés ({', '.join(ids_statuts)})")


def config_interface(cfg: dict) -> dict | None:
    """Paramètres utiles à l'interface (rayons, rôles de zone, types, catégories)."""
    c = cfg.get("crise")
    if not c:
        return None
    return {
        "rayon_max_m": c["rayon_max_m"], "rayon_defaut_m": c["rayon_defaut_m"], "pas_m": c["pas_m"],
        "raccourcis_m": c["raccourcis_m"], "types": c["types"], "zones": c["zones"],
        "categories": [{"id": k["id"], "libelle": k["libelle"], "role": k.get("role", "points"),
                        "propositions": k.get("propositions", {}), "couleur": k.get("couleur"), "nombre": k.get("nombre"),
                        "statuts": k.get("statuts") or [x["id"] for x in statuts(cfg)]} for k in c["categories"]],
        "statuts": statuts(cfg),
        "auteur_defaut": auteur_poste(),
        "vent": bool(c.get("vent")), "adresses": bool(c.get("geocodage")),
        "tampon_couloir_m": c.get("tampon_couloir_m", 100),
        "avertissement": c.get("avertissement", ""),
    }


def statuts(cfg: dict) -> list[dict]:
    return cfg["crise"].get("statuts") or STATUTS_DEFAUT


def auteur_poste() -> str:
    """Nom de session Windows du poste : auteur par défaut des décisions (modifiable dans l'interface)."""
    try:
        return getpass.getuser()
    except Exception:   # session sans nom (service) : l'utilisateur saisit le sien
        return ""


# ---------------------------------------------------------------- géométrie

def vers_l93(lon: float, lat: float) -> tuple[float, float]:
    x, y = _VERS_L93.transform(lon, lat)
    if not (math.isfinite(x) and math.isfinite(y)):
        raise ErreurCrise("point hors du domaine de Lambert 93")
    return x, y


def vers_wgs84(x: float, y: float) -> list[float]:
    lon, lat = _VERS_WGS84.transform(x, y)
    return [round(lon, DECIMALES_WGS84), round(lat, DECIMALES_WGS84)]


def geojson_wgs84(geom_l93, tolerance_m: float = 1.0) -> dict | None:
    """Géométrie Lambert 93 → GeoJSON WGS84, simplifiée (dessin seulement)."""
    if geom_l93 is None or geom_l93.is_empty:
        return None
    g = geom_l93.simplify(tolerance_m) if tolerance_m else geom_l93
    g = shapely.transform(g, lambda xy: np.column_stack(_VERS_WGS84.transform(xy[:, 0], xy[:, 1])))
    return mapping(shapely.set_precision(g, 10 ** -DECIMALES_WGS84))


def direction(dx: float, dy: float) -> str:
    """Direction (8 secteurs) d'un vecteur Lambert 93 (Y vers le nord)."""
    if dx == 0 and dy == 0:
        return "—"
    angle = (math.degrees(math.atan2(dx, dy)) + 360) % 360
    return DIRECTIONS[int((angle + 22.5) // 45) % 8]


def zones_geometriques(zones: list[dict], centre_l93: tuple[float, float], roles: list[str]) -> list[dict]:
    """Zones de la demande, contrôlées, avec leur géométrie Lambert 93 et leur rang de gravité."""
    sortie = []
    for i, z in enumerate(zones, 1):
        role = z.get("role")
        if role not in roles:
            raise ErreurCrise(f"zone n° {i} : rôle « {role} » inconnu (attendu : {', '.join(roles)})")
        forme = z.get("forme", "cercle")
        if forme == "cercle":
            r = z.get("rayon_m")
            if not isinstance(r, (int, float)) or isinstance(r, bool) or not r > 0:
                raise ErreurCrise(f"zone n° {i} : « rayon_m » doit être un nombre strictement positif")
            geom = Point(centre_l93).buffer(float(r), quad_segs=SEGMENTS_CERCLE)
            sortie.append({"role": role, "rang": roles.index(role), "forme": "cercle", "rayon_m": float(r), "geom": geom})
        elif forme == "polygone":
            pts = z.get("coordonnees")
            if not (isinstance(pts, list) and len(pts) >= 3):
                raise ErreurCrise(f"zone n° {i} : polygone d'au moins 3 sommets [lon, lat] attendu")
            try:
                geom = Polygon([vers_l93(float(p[0]), float(p[1])) for p in pts])
            except (TypeError, ValueError, IndexError) as e:
                raise ErreurCrise(f"zone n° {i} : sommets [lon, lat] attendus") from e
            if not geom.is_valid:
                geom = shapely.make_valid(geom)
            if geom.area <= 0:
                raise ErreurCrise(f"zone n° {i} : polygone de surface nulle")
            sortie.append({"role": role, "rang": roles.index(role), "forme": "polygone", "rayon_m": None, "geom": geom})
        elif forme == "couloir":
            # le long d'une ligne (canalisation, voie, cortège) : tampon de part et d'autre de l'axe (lot C3)
            pts, t = z.get("coordonnees"), z.get("tampon_m")
            if not (isinstance(pts, list) and len(pts) >= 2):
                raise ErreurCrise(f"zone n° {i} : couloir : axe d'au moins 2 sommets [lon, lat] attendu")
            if not isinstance(t, (int, float)) or isinstance(t, bool) or not 0 < t <= TAMPON_MAX_M:
                raise ErreurCrise(f"zone n° {i} : couloir : « tampon_m » attendu entre 0 et {TAMPON_MAX_M} m")
            try:
                axe = LineString([vers_l93(float(p[0]), float(p[1])) for p in pts])
            except (TypeError, ValueError, IndexError) as e:
                raise ErreurCrise(f"zone n° {i} : sommets [lon, lat] attendus") from e
            if axe.length <= 0:
                raise ErreurCrise(f"zone n° {i} : couloir de longueur nulle")
            geom = axe.buffer(float(t), quad_segs=SEGMENTS_CERCLE // 4)
            sortie.append({"role": role, "rang": roles.index(role), "forme": "couloir", "rayon_m": None, "geom": geom,
                           "tampon_m": float(t), "longueur_m": axe.length})
        else:
            raise ErreurCrise(f"zone n° {i} : forme « {forme} » inconnue (cercle, polygone ou couloir)")
    return sortie


def zone_de(geom, distance_m: float, zones: list[dict]) -> dict | None:
    """Zone la plus grave qui contient (ou touche) l'objet ; cercle : distance au centre ≤ rayon (exact)."""
    meilleure = None
    for z in zones:
        dedans = distance_m <= z["rayon_m"] if z["forme"] == "cercle" else geom.intersects(z["geom"])
        if dedans and (meilleure is None or z["rang"] < meilleure["rang"]):
            meilleure = z
    return meilleure


# ---------------------------------------------------------------- couches

def _couches(cfg: dict, libelles: set[str]) -> tuple[dict[str, inv_sig.Couche], dict[str, str]]:
    """Seules les couches demandées, par libellé : services (cache, sinon téléchargement), puis fichiers."""
    from . import esri as esri_sig
    services = {}
    for s, c in esri_sig.services_declares(cfg):
        services.setdefault(s["libelle"], (s, c))
    trouvees, anomalies = {}, {}
    for lib in sorted(libelles & services.keys()):
        s, c = services[lib]
        try:
            trouvees[lib] = esri_sig.couche_service(cfg, s, rang=0, connecteur=c)
        except inv_sig.AnomalieCouche as e:
            anomalies[lib] = str(e)
        except Exception as e:   # un service ne doit jamais interrompre l'analyse
            anomalies[lib] = f"erreur inattendue ({type(e).__name__} : {e})"
    reste = libelles - services.keys()
    if reste:
        for x in inv_sig.inventaire(cfg).couches:
            if x.libelle in reste:
                trouvees.setdefault(x.libelle, x)
    return trouvees, anomalies


def _version(couche: inv_sig.Couche) -> str:
    try:
        return str(couche._chemin.stat().st_mtime_ns)
    except (OSError, AttributeError):
        return "?"


def _objets_l93(couche: inv_sig.Couche):
    """(GeoDataFrame d'origine, géométries Lambert 93) gardés tant que la couche ne change pas."""
    cle = (couche.id, _version(couche))
    with _verrou:
        if cle in _objets_couches:
            return _objets_couches[cle]
    charge = obj_sig.charger(couche)
    geoms = charge.l93.geometry.values
    val = (charge.origine, geoms, shapely.STRtree(geoms))
    with _verrou:
        for k in [k for k in _objets_couches if k[0] == couche.id]:
            del _objets_couches[k]
        _objets_couches[cle] = val
    return val


def _source(couche: inv_sig.Couche) -> dict:
    s = couche.service or {}
    date = s.get("date_maj")
    if not date and couche._chemin is not None:
        try:
            date = datetime.fromtimestamp(couche._chemin.stat().st_mtime).isoformat(timespec="seconds")
        except OSError:
            date = None
    return {"couche": couche.id, "libelle": couche.libelle, "date": date, "origine": s.get("source", "fichier"),
            "avertissements": list(couche.avertissements[:3])}


def _texte(v) -> str | None:
    v = obj_sig.valeur_json(v)
    if v is None:
        return None
    t = str(v).strip()
    return t or None


# ---------------------------------------------------------------- candidats (objets à moins de rayon_max du point)

def _candidats_couche(couche, categorie: dict, centre: Point, rayon_max: float) -> list[dict]:
    return _candidats_objets(_objets_l93(couche), couche.id, couche.libelle, couche.champ_libelle, categorie, centre, rayon_max)


def _candidats_objets(objets, id_couche, libelle_couche, champ_libelle, categorie: dict, centre: Point,
                      rayon_max: float) -> list[dict]:
    """Objets à moins de `rayon_max` du point, avec leur distance et les champs utiles de la catégorie."""
    origine, geoms, arbre = objets
    proches = arbre.query(centre.buffer(rayon_max), predicate="intersects")
    if not len(proches):
        return []
    dist = shapely.distance(geoms[proches], centre)
    regroupement = categorie.get("regroupement")

    def colonne(nom):
        """Valeurs du champ pour les objets proches (lecture en bloc : des dizaines de milliers d'objets) ;
        une liste de champs donne la première valeur renseignée (nom de voie à gauche, sinon à droite)."""
        noms = [x for x in (nom if isinstance(nom, list) else [nom]) if x and x in origine.columns]
        vals = [None] * len(proches)
        for x in noms:
            col = origine[x].to_numpy()[proches]
            vals = [v if v is not None else _texte(w) for v, w in zip(vals, col)]
        return vals

    libelles = colonne(categorie.get("champ_libelle") or champ_libelle)
    # « @couche » : un seul élément par couche (réseau de gaz, dont les tronçons n'ont pas de nom)
    groupes = [libelle_couche] * len(proches) if regroupement == "@couche" else colonne(regroupement)
    details, types = colonne(categorie.get("champ_detail")), colonne(categorie.get("champ_type"))
    ordres = colonne(categorie.get("champ_ordre"))
    sans_nom = categorie.get("libelle_sans_nom")
    if sans_nom and regroupement and regroupement != "@couche":
        # voies sans nom regroupées par nature (« Voie sans nom (Rond-point) »), jamais une par tronçon
        groupes = [g or (f"{sans_nom} ({t})" if t else sans_nom) for g, t in zip(groupes, types)]
    numeros = origine.index.to_numpy()[proches]
    ponctuels = shapely.get_type_id(geoms[proches]) == 0
    sortie = []
    for i, (n, d) in enumerate(zip(proches.tolist(), dist.tolist())):
        g = geoms[n]
        sortie.append({
            "n": int(numeros[i]), "couche": id_couche, "geom": g, "distance_m": float(d),
            "proche": g if ponctuels[i] else None,   # calculé pour l'élément retenu seulement (_element)
            "libelle": libelles[i], "groupe": groupes[i], "detail": details[i], "type": types[i], "ordre": ordres[i],
        })
    return sortie


def _ouvrage(fichier: str, dossiers: dict) -> tuple[str, str, str] | None:
    """(clé de l'ouvrage, nom, type) d'une maquette selon son chemin : dossier qui suit le dossier de type."""
    parties = fichier.replace("\\", "/").split("/")
    types = {k.lower(): v for k, v in dossiers.items()}
    for i, p in enumerate(parties[:-1]):
        if p.lower() in types:
            nom = parties[i + 1] if i + 1 < len(parties) - 1 else Path(parties[-1]).stem
            return "/".join(parties[:i + 2]), nom.strip(), p.lower()
    return None


def _candidats_maquettes(cfg: dict, categorie: dict, centre: Point, rayon_max: float) -> tuple[list[dict], list[str]]:
    from . import maquettes as maq_sig
    try:
        cfgm = maq_sig.config()
    except maq_sig.ConfigMaquettesInvalide as e:
        return [], [f"maquettes : {e}"]
    # liste déjà tenue à jour par l'interface : pas de rebalayage du dossier (des secondes au démarrage du
    # serveur, verrou partagé avec /maquettes) ; seulement si elle n'a encore jamais été construite
    if not maq_sig.suivi.lignes:
        maq_sig.suivi.actualiser(cfg, cfgm)
    o, s = vers_wgs84(centre.x - rayon_max, centre.y - rayon_max)
    e, n = vers_wgs84(centre.x + rayon_max, centre.y + rayon_max)
    sortie, anomalies = [], []
    for ligne in maq_sig.suivi.dans_emprise([o - 0.01, s - 0.01, e + 0.01, n + 0.01]):
        ouvrage = _ouvrage(ligne.get("fichier", ""), categorie["dossiers"])
        if ouvrage is None:
            continue
        m = maq_sig.une_maquette(cfg, cfgm, ligne["id"])
        pos = (m or {}).get("positionnement") or {}
        if not pos.get("emprise_geojson"):
            continue
        try:
            from .croisement import geometrie_l93
            g = geometrie_l93(pos["emprise_geojson"])
        except Exception as ex:   # emprise illisible : la maquette est signalée, les autres continuent
            anomalies.append(f"{ligne.get('fichier')} : emprise illisible ({type(ex).__name__})")
            continue
        d = float(g.distance(centre))
        if d > rayon_max:
            continue
        sortie.append({"n": None, "couche": None, "geom": g, "distance_m": d, "proche": nearest_points(g, centre)[0],
                       "libelle": ouvrage[1], "groupe": ouvrage[0], "type": ouvrage[2], "detail": None,
                       "maquette": ligne["id"]})
    return sortie, anomalies


def candidats(cfg: dict, centre_l93: tuple[float, float]) -> dict:
    """Objets de chaque catégorie à moins de « rayon_max_m » du point (mis en cache par point)."""
    c = cfg["crise"]
    rayon_max = float(c["rayon_max_m"])
    besoins = {lib for k in c["categories"] for lib in k.get("couches", [])}
    par_libelle, anomalies_couches = _couches(cfg, besoins)
    versions = tuple(sorted((lib, par_libelle[lib].id, _version(par_libelle[lib])) for lib in besoins if lib in par_libelle))
    cle = (round(centre_l93[0], 1), round(centre_l93[1], 1), rayon_max, versions,
           json.dumps(c["categories"], sort_keys=True, ensure_ascii=False))
    with _verrou:
        if cle in _candidats:
            _candidats.move_to_end(cle)
            return _candidats[cle]
    centre = Point(centre_l93)
    resultat = {"categories": {}, "sources": [], "anomalies": []}
    for k in c["categories"]:
        role = k.get("role", "points")
        objets, sources = [], []
        if role == "maquettes":
            objets, anom = _candidats_maquettes(cfg, k, centre, rayon_max)
            resultat["anomalies"] += [f"{k['libelle']} : {a}" for a in anom]
            sources.append({"couche": None, "libelle": "Maquettes IFC positionnées", "date": None, "origine": "maquettes",
                            "avertissements": []})
        for lib in k.get("couches", []):
            couche = par_libelle.get(lib)
            if couche is None:
                motif = anomalies_couches.get(lib)
                resultat["anomalies"].append(f"{k['libelle']} : couche « {lib} » indisponible"
                                             + (f" ({motif})" if motif else " (non déclarée ou illisible)"))
                continue
            try:
                # secours : cherchés plus loin (le commissariat le plus proche peut être au-delà de rayon_max)
                rayon = float(k.get("rayon_recherche_m", rayon_max)) if role == "secours" else rayon_max
                objets += _candidats_couche(couche, k, centre, rayon)
                sources.append(_source(couche))
            except inv_sig.AnomalieCouche as e:
                resultat["anomalies"].append(f"{k['libelle']} : {lib} : {e}")
        resultat["categories"][k["id"]] = objets
        resultat["sources"] += [{**s, "categorie": k["id"]} for s in sources]
    with _verrou:
        _candidats[cle] = resultat
        while len(_candidats) > CACHE_CANDIDATS:
            _candidats.popitem(last=False)
    return resultat


# ---------------------------------------------------------------- analyse

def _element(cat: dict, cle: str, libelle: str, membres: list[dict], centre: Point, zones: list[dict], **extra) -> dict:
    """Élément de la liste : membre le plus proche, zone la plus grave parmi celles de ses membres."""
    proche = min(membres, key=lambda m: m["distance_m"])
    meilleure = None
    # à gravité égale, le point de la station l'emporte sur ses accès (la cause « par l'accès » n'est dite que
    # si la station elle-même est hors de cette zone), puis le membre le plus proche
    for m in sorted(membres, key=lambda m: (bool(m.get("acces")), m["distance_m"])):
        z = zone_de(m["geom"], m["distance_m"], zones)
        if z is not None and (meilleure is None or z["rang"] < meilleure[0]["rang"]):
            meilleure = (z, m)
    p = proche["proche"] if proche.get("proche") is not None else nearest_points(proche["geom"], centre)[0]
    e = {
        "categorie": cat["id"], "cle": cle, "libelle": libelle,
        "distance_m": round(proche["distance_m"], 1),
        "direction": direction(p.x - centre.x, p.y - centre.y),
        "position": (p.x, p.y),   # Lambert 93 ; converti en WGS84 par lot dans analyser()
        "zone": meilleure[0]["role"] if meilleure else None,
        "proposition": (cat.get("propositions") or {}).get(meilleure[0]["role"], "") if meilleure else "",
        "couche": proche["couche"], "n": proche["n"],
    }
    e.update(extra)
    if meilleure:
        e["_membre_zone"] = meilleure[1]
    return e


def _detail(membres: list[dict]) -> str | None:
    vals = sorted({m["detail"] for m in membres if m.get("detail")})
    return ", ".join(vals) if vals else None


def _groupes(objets: list[dict]) -> dict[str, list[dict]]:
    g: dict[str, list[dict]] = {}
    for o in objets:
        cle = o["groupe"] or f"{o['couche']}:{o['n']}"
        g.setdefault(cle, []).append(o)
    return g


def _analyser_points(cat, objets, centre, zones):
    return [_element(cat, cle, m[0]["groupe"] or m[0]["libelle"] or f"objet n° {m[0]['n']}", m, centre, zones,
                     detail=_detail(m))
            for cle, m in _groupes(objets).items()]


def _analyser_stations(cat, objets, acces, cfg_crise, centre, zones):
    """Stations (regroupées) touchées par leur point ou par l'un de leurs accès (rattachés au plus proche)."""
    groupes = _groupes(objets)
    rayon = float(cfg_crise["rattachement_acces_m"])
    if groupes and acces:
        noms = list(groupes)
        pts = [unary_union([m["geom"] for m in groupes[n]]) for n in noms]
        arbre = shapely.STRtree(pts)
        for a in acces:
            i = arbre.nearest(a["geom"])
            if i is not None and pts[i].distance(a["geom"]) <= rayon:
                a["station"] = noms[int(i)]
    elements = []
    for cle, membres in groupes.items():
        lies = [a for a in acces if a.get("station") == cle]
        e = _element(cat, cle, cle, membres + [{**a, "acces": True} for a in lies], centre, zones,
                     detail=_detail(membres), nb_acces=len(lies))
        m = e.get("_membre_zone")
        if m is not None and m.get("acces"):
            # touchée seulement par un accès : la cause est dite (le point de la station est hors de cette zone)
            e["cause"] = f"par l'accès « {m['libelle'] or 'sans nom'} » ({round(m['distance_m'])} m)"
        elements.append(e)
    return elements


def _analyser_acces(cat, acces, centre, zones):
    sortie = []
    for a in acces:
        e = _element(cat, f"{a['couche']}:{a['n']}", a["libelle"] or "Accès sans nom", [a], centre, zones,
                     detail=f"station {a['station']} (rattachement par proximité)" if a.get("station")
                     else "aucune station à proximité")
        sortie.append(e)
    return sortie


def _analyser_lignes(cat, objets, centre, zones):
    union = unary_union([z["geom"] for z in zones]) if zones else None
    sortie = []
    for cle, membres in _groupes(objets).items():
        e = _element(cat, cle, membres[0]["groupe"] or membres[0]["libelle"] or cle, membres, centre, zones)
        if e["zone"] and union is not None:
            zone = next(z for z in zones if z["role"] == e["zone"])
            trace = unary_union([m["geom"] for m in membres])
            e["longueur_zone_m"] = round(trace.intersection(zone["geom"]).length) if zone["forme"] != "cercle" else \
                round(trace.intersection(Point(centre).buffer(zone["rayon_m"], quad_segs=SEGMENTS_CERCLE)).length)
            e["geometrie"] = geojson_wgs84(trace.intersection(union))
            e["detail"] = f"{e['longueur_zone_m']:,} m de voie dans la zone".replace(",", " ")
        sortie.append(e)
    return sortie


def _analyser_plai(cat, objets, centre, zones):
    bloquantes = [z for z in zones if z["role"] in cat.get("zones_bloquantes", [])]
    rassemblement = cat["type_rassemblement"]
    sortie = []
    for o in objets:
        e = _element(cat, f"{o['couche']}:{o['n']}", o["libelle"] or o["type"] or "Point PLAI", [o], centre, zones,
                     type=o["type"], detail=o["detail"])
        sortie.append(e)
    libres = [o for o in objets if o["type"] == rassemblement
              and not any(zone_de(o["geom"], o["distance_m"], [z]) for z in bloquantes)]
    for e, o in zip(sortie, objets):
        if o["type"] != rassemblement or not e["zone"] or e["zone"] not in cat.get("zones_bloquantes", []):
            continue
        e["inutilisable"] = True
        if libres:
            r = min(libres, key=lambda x: x["geom"].distance(o["geom"]))
            e["remplacement"] = {"libelle": r["libelle"] or "point de rassemblement", "couche": r["couche"], "n": r["n"],
                                 "position": vers_wgs84(r["proche"].x, r["proche"].y),
                                 "distance_m": round(r["geom"].distance(o["geom"]))}
            e["proposition"] = (f"Inutilisable : utiliser « {e['remplacement']['libelle']} » "
                                f"à {e['remplacement']['distance_m']} m, hors zone")
        else:
            e["proposition"] = "Inutilisable : aucun point de rassemblement connu hors zone, à définir sur place"
    return sortie


def _analyser_secours(cat, objets, centre, zones):
    """Les « nombre » sites les plus proches du point (commissariats, casernes), dans ou hors des zones."""
    sortie = []
    for i, o in enumerate(sorted(objets, key=lambda o: o["distance_m"])[:int(cat.get("nombre", 1))], 1):
        e = _element(cat, f"{o['couche']}:{o['n']}", o["libelle"] or o["detail"] or "Site sans nom", [o], centre, zones,
                     detail=o["detail"], secours=True, rang=i)
        if not e["zone"]:
            e["proposition"] = "Hors zone" + (" : le plus proche" if i == 1 else "")
        sortie.append(e)
    return sortie


def _analyser_maquettes(cat, objets, centre, zones):
    puits = {p.lower() for p in cat.get("puits", [])}
    types = {k.lower(): v for k, v in cat["dossiers"].items()}
    sortie = []
    for cle, membres in _groupes(objets).items():
        t = membres[0]["type"]
        detail = f"{types.get(t, t)} · {len(membres)} maquette(s)"
        e = _element(cat, cle, membres[0]["libelle"], membres, centre, zones, detail=detail, puits=t in puits,
                     maquettes=sorted({m["maquette"] for m in membres}))
        if e["puits"] and e["zone"]:
            e["proposition"] = (e["proposition"] + " — " if e["proposition"] else "") + cat.get(
                "message_puits", "puits : fumées ou gaz peuvent passer vers le souterrain")
        if e["zone"]:
            e["geometrie"] = geojson_wgs84(unary_union([m["geom"] for m in membres]), 0.5)
        sortie.append(e)
    return sortie


def _nombre(v) -> float | None:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def _analyser_bouclage(cat, objets, centre, zones):
    """Points où les voies coupent la limite extérieure des zones bouclées : un point à barrer par voie et par
    passage (points d'une même voie à moins de FUSION_BOUCLAGE_M fusionnés). Zone fixée par le serveur : le point
    est sur la limite, un recalcul par la distance au centre le ferait basculer d'un côté ou de l'autre."""
    zb = [z for z in zones if z["role"] in cat["zones_bouclage"]]
    if not zb or not objets:
        return []
    limite = unary_union([z["geom"] for z in zb]).boundary
    fusion = float(cat.get("fusion_m", FUSION_BOUCLAGE_M))
    points: dict[str, list[tuple[Point, dict]]] = {}
    for o in objets:
        if not o["geom"].intersects(limite):
            continue
        coupe = o["geom"].intersection(limite)
        for p in getattr(coupe, "geoms", [coupe]):
            if p.is_empty:
                continue
            p = p if p.geom_type == "Point" else p.representative_point()   # tronçon qui longe la limite
            voie = o["groupe"] or o["libelle"] or "Voie sans nom"
            deja = points.setdefault(voie, [])
            if not any(q.distance(p) < fusion for q, _ in deja):
                deja.append((p, o))
    sortie = []
    for voie, liste in points.items():
        for p, o in liste:
            z = min(zb, key=lambda z: z["geom"].boundary.distance(p))
            ordre = _nombre(o.get("ordre"))
            detail = ", ".join(x for x in (o.get("type"), f"importance {o['ordre']}" if o.get("ordre") else None) if x)
            m = {**o, "geom": p, "proche": p, "distance_m": float(p.distance(centre))}
            e = _element(cat, f"{voie}@{round(p.x)},{round(p.y)}", voie, [m], centre, zones, detail=detail or None,
                         fixe=True, ordre=ordre)
            e["zone"] = z["role"]
            e["proposition"] = (cat.get("propositions") or {}).get(z["role"], "Barrer la voie")
            sortie.append(e)
    return sortie


def _regions_exclusives(zones: list[dict]) -> list[tuple[dict, object]]:
    """Chaque zone privée des zones plus graves : un objet compte dans la plus grave qui le contient."""
    sortie, deja = [], None
    for z in sorted(zones, key=lambda z: z["rang"]):
        sortie.append((z, z["geom"] if deja is None else z["geom"].difference(deja)))
        deja = z["geom"] if deja is None else deja.union(z["geom"])
    return sortie


def _analyser_denombrement(cat, objets, centre, zones, rayon_donnees: float):
    """Mesures (nombre d'objets, somme d'un champ) par zone et courbe « valeur dans un cercle de rayon r ».

    Surfaces au prorata (« prorata: true », carreaux de population) : chaque objet compte pour la part de sa
    surface dans la zone ; sinon un objet compte entier dans la zone la plus grave qu'il touche (bâtiments).
    La courbe sert à l'interface pendant le glissement ; les valeurs par zone sont exactes (polygones compris)."""
    pas = float(cat.get("pas_courbe_m", PAS_COURBE_M))
    rayons = np.arange(0, rayon_donnees + pas, pas)
    prorata = bool(cat.get("prorata"))
    geoms = np.array([o["geom"] for o in objets], dtype=object)
    dist = np.array([o["distance_m"] for o in objets], dtype=float)
    valeurs = {}
    for m in cat["mesures"]:
        champ = m.get("champ")
        v = np.array([1.0 if not champ else (_nombre(o["valeurs"].get(champ)) or 0.0) for o in objets], dtype=float)
        valeurs[m["id"]] = v
    if prorata and len(objets):
        surfaces = shapely.area(geoms)
    cle = (cat["id"], round(centre.x, 1), round(centre.y, 1), pas, len(objets), id(objets[0]["geom"]) if objets else 0)
    with _verrou:
        parts = _parts_courbes.get(cle)
    if prorata and len(objets) and parts is None:   # ne dépend que du point : gardé pendant le glissement
        loin = shapely.hausdorff_distance(geoms, centre)   # sommet le plus éloigné : au-delà, objet entier dans le cercle
        parts = np.zeros((len(rayons), len(objets)))
        for i, r in enumerate(rayons):
            dedans = loin <= r
            parts[i, dedans] = 1.0
            bord = (~dedans) & (dist < r)
            if bord.any():
                disque = centre.buffer(float(r), quad_segs=SEGMENTS_CERCLE)
                parts[i, bord] = shapely.area(shapely.intersection(geoms[bord], disque)) / np.where(surfaces[bord] > 0, surfaces[bord], 1)
        with _verrou:
            _parts_courbes[cle] = parts
            while len(_parts_courbes) > CACHE_CANDIDATS:
                _parts_courbes.pop(next(iter(_parts_courbes)))
    mesures = []
    regions = _regions_exclusives(zones)
    for m in cat["mesures"]:
        v = valeurs[m["id"]]
        if prorata and len(objets):
            courbe = parts @ v
        else:
            ordre = np.argsort(dist)
            cumul = np.concatenate([[0.0], np.cumsum(v[ordre])])
            courbe = cumul[np.searchsorted(dist[ordre], rayons, side="right")]
        par_zone = {}
        for z, region in regions:
            if prorata and len(objets):
                touches = shapely.intersects(geoms, region)
                inter = np.zeros(len(objets))
                if touches.any():
                    inter[touches] = shapely.area(shapely.intersection(geoms[touches], region)) / np.where(
                        surfaces[touches] > 0, surfaces[touches], 1)
                par_zone[z["role"]] = float(inter @ v) + par_zone.get(z["role"], 0.0)
            else:
                par_zone.setdefault(z["role"], 0.0)
        if not prorata:
            for o, x in zip(objets, v):
                z = zone_de(o["geom"], o["distance_m"], zones)
                if z is not None:
                    par_zone[z["role"]] = par_zone.get(z["role"], 0.0) + float(x)
        mesures.append({"id": m["id"], "libelle": m["libelle"], "arrondi": m.get("arrondi", 1),
                        "par_zone": {k: round(x, 1) for k, x in par_zone.items()},
                        "courbe": np.round(courbe, 1).tolist()})
    return {"mesures": mesures, "pas_courbe_m": pas, "rayon_donnees_m": rayon_donnees,
            "date_donnee": cat.get("date_donnee"), "objets": len(objets)}


def distance_texte(m: float) -> str:
    return f"{m / 1000:g} km".replace(".", ",") if m >= 1000 else f"{m:g} m"


def _hors_donnees(zones: list[dict], centre: Point, rayon: float) -> list[str]:
    """Zones qui dépassent le disque des données d'une source locale (résultats incomplets au-delà)."""
    disque = centre.buffer(rayon + 1, quad_segs=SEGMENTS_CERCLE)
    return [f"{z['role']} ({round(z['rayon_m'])} m)" if z["forme"] == "cercle" else z["role"]
            for z in zones if not disque.contains(z["geom"])]


def _locaux(cfg: dict, k: dict, centre: Point, etats: dict, fabrique) -> tuple[list[dict], dict]:
    """Objets d'une catégorie à source locale (rues, bâtiments, population) et état de la source (lue une fois
    par analyse : rues et points de bouclage partagent la même source)."""
    s = crise_local.source(cfg, k["source_locale"])
    if s["id"] not in etats:
        etats[s["id"]] = crise_local.objets(cfg, s, (centre.x, centre.y), **({"fabrique": fabrique} if fabrique else {}))
    etat = etats[s["id"]]
    if etat["objets"] is None:
        return [], {**etat, "source": s}
    objets = _candidats_objets(etat["objets"], None, s["libelle"], None, k, centre, float(s["rayon_m"]))
    if k.get("role") == "denombrement" and objets:
        origine = etat["objets"][0]
        n = np.array([o["n"] for o in objets])
        colonnes = {m["champ"]: origine[m["champ"]].to_numpy()[n]
                    for m in k["mesures"] if m.get("champ") and m["champ"] in origine.columns}
        for i, o in enumerate(objets):
            o["valeurs"] = {c: v[i] for c, v in colonnes.items()}
    return objets, {**etat, "source": s}


def analyser(cfg: dict, centre_wgs84: list[float], zones: list[dict], fabrique=None) -> dict:
    """Éléments touchés par catégorie et par zone, plus tous les candidats jusqu'à « rayon_max_m »."""
    c = cfg["crise"]
    try:
        lon, lat = float(centre_wgs84[0]), float(centre_wgs84[1])
    except (TypeError, ValueError, IndexError) as e:
        raise ErreurCrise("« centre » attendu : [longitude, latitude]") from e
    centre_l93 = vers_l93(lon, lat)
    roles = [z["id"] for z in c["zones"]]
    zg = zones_geometriques(zones or [], centre_l93, roles)
    base = candidats(cfg, centre_l93)
    centre = Point(centre_l93)
    cats = {k["id"]: k for k in c["categories"]}
    acces_cat = next((k for k in c["categories"] if k.get("role") == "acces"), None)
    acces = [dict(a) for a in base["categories"].get(acces_cat["id"], [])] if acces_cat else []
    categories = []
    anomalies = list(base["anomalies"])   # copie : la base est gardée en cache
    sources = list(base["sources"])
    chargement = []
    etats: dict[str, dict] = {}
    for k in c["categories"]:
        role = k.get("role", "points")
        objets = base["categories"].get(k["id"], [])
        extra = {}
        if k.get("source_locale"):
            id_s = k["source_locale"]
            deja = id_s in etats
            objets, etat = _locaux(cfg, k, centre, etats, fabrique)
            s = etat["source"]
            if not deja:   # état de la source signalé une fois (rues et points de bouclage la partagent)
                anomalies += [f"{s['libelle']} : {a}" for a in etat["anomalies"]]
                if etat["en_cours"]:
                    chargement.append({"source": s["libelle"], "progression": etat["progression"]})
                elif etat["objets"] is None:
                    anomalies.append(f"{s['libelle']} : service injoignable et aucun cache autour de ce point")
                else:
                    sources.append({"couche": None, "libelle": s["libelle"], "date": s.get("date_donnee") or etat["date"],
                                    "origine": etat["origine"], "avertissements": [], "categorie": k["id"]})
                    dehors = _hors_donnees(zg, centre, float(s["rayon_m"]))
                    if dehors:
                        anomalies.append(f"{s['libelle']} : données chargées jusqu'à {distance_texte(s['rayon_m'])} du point ; "
                                         f"zone(s) au-delà : {', '.join(dehors)} (résultats incomplets)")
            if role == "denombrement" and etat["objets"] is not None:
                manquants = [m["champ"] for m in k["mesures"] if m.get("champ") and m["champ"] not in etat["objets"][0].columns]
                if manquants and len(etat["objets"][0]):
                    anomalies.append(f"{k['libelle']} : champ(s) absent(s) de la source : {', '.join(manquants)}")
            extra = {"en_cours": etat["en_cours"], "source": s["libelle"]}
        if role == "denombrement":
            d = _analyser_denombrement(k, objets, centre, zg, float(crise_local.source(cfg, k["source_locale"])["rayon_m"])) \
                if not extra.get("en_cours") else {"mesures": []}
            categories.append({"id": k["id"], "libelle": k["libelle"], "role": role, "elements": [], "touches": 0, **extra, **d})
            continue
        if role == "bouclage":
            elements = _analyser_bouclage(k, objets, centre, zg)
        elif role in ("lignes", "rues"):
            elements = _analyser_lignes(k, objets, centre, zg)
        elif role == "stations":
            elements = _analyser_stations(k, objets, acces, c, centre, zg)
        elif role == "acces":
            elements = _analyser_acces(k, acces, centre, zg)
        elif role == "plai":
            elements = _analyser_plai(k, objets, centre, zg)
        elif role == "maquettes":
            elements = _analyser_maquettes(k, objets, centre, zg)
        elif role == "secours":
            elements = _analyser_secours(k, objets, centre, zg)
            if not elements:   # rien dans le rayon de recherche : signalé, jamais ignoré en silence
                anomalies.append(f"{k['libelle']} : aucun site à moins de {k.get('rayon_recherche_m', c['rayon_max_m']) / 1000:g} km")
        else:
            elements = _analyser_points(k, objets, centre, zg)
        for e in elements:
            e.pop("_membre_zone", None)
        if elements:   # positions converties en WGS84 en un seul appel (milliers d'objets)
            xy = np.array([e["position"] for e in elements])
            lon_, lat_ = _VERS_WGS84.transform(xy[:, 0], xy[:, 1])
            for e, a, b in zip(elements, np.round(lon_, DECIMALES_WGS84).tolist(), np.round(lat_, DECIMALES_WGS84).tolist()):
                e["position"] = [a, b]
        if role == "bouclage":   # voies principales d'abord (importance 1 à 5 de la BD TOPO), puis la distance
            elements.sort(key=lambda e: (roles.index(e["zone"]), e["ordre"] if e.get("ordre") is not None else 99, e["distance_m"]))
        else:
            elements.sort(key=lambda e: (roles.index(e["zone"]) if e["zone"] else len(roles), e["distance_m"]))
        categories.append({"id": k["id"], "libelle": k["libelle"], "role": role, "elements": elements,
                           "touches": sum(1 for e in elements if e["zone"]), **extra})
    return {
        "centre": {"wgs84": [lon, lat], "l93": [round(centre_l93[0], 2), round(centre_l93[1], 2)]},
        "rayon_max_m": c["rayon_max_m"],
        "zones": [{"role": z["role"], "forme": z["forme"], "rayon_m": z["rayon_m"], "surface_m2": round(z["geom"].area),
                   # couloir : dessiné par l'interface à partir de la géométrie calculée ici (tampon en Lambert 93)
                   **({"tampon_m": z["tampon_m"], "longueur_m": round(z["longueur_m"]), "geometrie": geojson_wgs84(z["geom"], 0.5)}
                      if z["forme"] == "couloir" else {})}
                  for z in zg],
        "categories": categories,
        "sources": sources,
        "chargement": chargement,
        "anomalies": anomalies,
        "date": datetime.now().isoformat(timespec="seconds"),
    }


# ---------------------------------------------------------------- recherche d'un lieu

def lieux(cfg: dict, q: str, max_resultats: int = 15) -> list[dict]:
    """Stations et accès des catégories « stations » et « acces » dont le nom contient le texte."""
    cible = obj_sig.normaliser(q.strip())
    if len(cible) < 2:
        return []
    besoins = {lib for k in cfg["crise"]["categories"] if k.get("role") in ("stations", "acces") for lib in k["couches"]}
    par_libelle, _ = _couches(cfg, besoins)
    sortie, vus = [], set()
    for k in cfg["crise"]["categories"]:
        role = k.get("role", "points")
        if role not in ("stations", "acces"):
            continue
        for lib in k.get("couches", []):
            couche = par_libelle.get(lib)
            if couche is None:
                continue
            try:
                origine, geoms, _ = _objets_l93(couche)
            except inv_sig.AnomalieCouche:
                continue
            champ = k.get("regroupement") if role == "stations" and k.get("regroupement") else (k.get("champ_libelle") or couche.champ_libelle)
            if not champ or champ not in origine.columns:
                continue
            noms = origine[champ].astype(str)
            for i in np.flatnonzero(noms.map(obj_sig.normaliser).str.contains(cible, regex=False).to_numpy()):
                nom = noms.iloc[i]
                cle = (role, nom)
                if role == "stations" and cle in vus:
                    continue
                vus.add(cle)
                p = geoms[i].centroid
                sortie.append({"type": "Station" if role == "stations" else "Accès", "libelle": nom,
                               "position": vers_wgs84(p.x, p.y)})
    sortie.sort(key=lambda s: (s["type"] != "Station", not obj_sig.normaliser(s["libelle"]).startswith(cible), s["libelle"]))
    return sortie[:max_resultats]


# ---------------------------------------------------------------- export CSV

COLONNES = [("zone", "Zone"), ("categorie", "Catégorie"), ("libelle", "Élément"), ("detail", "Détail"),
            ("cause", "Cause"), ("distance_m", "Distance (m)"), ("direction", "Direction"),
            ("proposition", "Proposition"), ("lon", "Longitude"), ("lat", "Latitude"), ("source", "Source"),
            ("date_source", "Date de la donnée")]


def lignes_csv(cfg: dict, r: dict) -> list[dict]:
    zones = {z["id"]: z["libelle"] for z in cfg["crise"]["zones"]}
    sources = {}
    for s in r["sources"]:
        sources.setdefault(s["categorie"], []).append(s)
    lignes = []
    for k in r["categories"]:
        src = sources.get(k["id"], [])
        for e in k["elements"]:
            if not e["zone"] and k["role"] != "secours":
                continue
            s = next((x for x in src if x["couche"] == e.get("couche")), src[0] if src else {})
            lignes.append({"zone": zones.get(e["zone"], e["zone"]) if e["zone"] else "hors zone", "categorie": k["libelle"], "libelle": e["libelle"],
                           "detail": e.get("detail"), "cause": e.get("cause"), "distance_m": e["distance_m"],
                           "direction": e["direction"], "proposition": e.get("proposition"),
                           "lon": e["position"][0], "lat": e["position"][1],
                           "source": s.get("libelle"), "date_source": s.get("date")})
    return lignes


# ---------------------------------------------------------------- événements enregistrés (poste local)

def dossier_evenements(cfg: dict) -> Path:
    from .config import RACINE
    p = Path(cfg["crise"]["dossier_travail"])
    return p if p.is_absolute() else RACINE / p


_ID = re.compile(r"^[0-9]{8}-[0-9]{6}$")


def enregistrer(cfg: dict, evt: dict) -> dict:
    """Écrit l'événement (description, centre, zones) dans le dossier de travail ; renvoie son identifiant."""
    if not isinstance(evt, dict):
        raise ErreurCrise("événement attendu")
    titre = evt.get("titre")
    if not isinstance(titre, str) or not titre.strip():
        raise ErreurCrise("« titre » de l'événement attendu")
    if evt.get("mode") not in ("reel", "exercice"):
        raise ErreurCrise("« mode » attendu : reel ou exercice")
    if evt.get("centre") is not None:
        vers_l93(*map(float, evt["centre"]))
        zones_geometriques(evt.get("zones") or [], vers_l93(*map(float, evt["centre"])),
                           [z["id"] for z in cfg["crise"]["zones"]])
    id_ = evt.get("id") or datetime.now().strftime("%Y%m%d-%H%M%S")
    if not _ID.match(id_):
        raise ErreurCrise("identifiant d'événement invalide")
    ids_statuts = {x["id"] for x in statuts(cfg)}
    decisions = evt.get("statuts") or {}
    if not isinstance(decisions, dict) or any(not isinstance(v, dict) or v.get("statut") not in ids_statuts
                                              for v in decisions.values()):
        raise ErreurCrise(f"« statuts » : statut inconnu (attendu : {', '.join(sorted(ids_statuts))})")
    journal = evt.get("journal") or []
    if not isinstance(journal, list) or any(not isinstance(j, dict) or not isinstance(j.get("texte"), str) for j in journal):
        raise ErreurCrise("« journal » : liste d'entrées {heure, auteur, type, texte} attendue")
    versions = evt.get("versions") or []
    if not isinstance(versions, list):
        raise ErreurCrise("« versions » : liste attendue")
    garde = {k: evt.get(k) for k in ("titre", "type", "mode", "debut", "description", "centre", "zones", "auteur")}
    garde.update({"statuts": decisions, "journal": journal[-MAX_JOURNAL:], "versions": versions,
                  "id": id_, "modifie": datetime.now().isoformat(timespec="seconds")})
    d = dossier_evenements(cfg)
    d.mkdir(parents=True, exist_ok=True)
    tmp = d / f"{id_}.json.tmp"
    tmp.write_text(json.dumps(garde, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(d / f"{id_}.json")
    return garde


def evenements(cfg: dict) -> list[dict]:
    d = dossier_evenements(cfg)
    sortie = []
    for p in sorted(d.glob("*.json"), reverse=True) if d.is_dir() else []:
        try:
            e = json.loads(p.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        sortie.append({k: e.get(k) for k in ("id", "titre", "type", "mode", "debut", "modifie")})
    return sortie


def evenement(cfg: dict, id_: str) -> dict | None:
    if not _ID.match(id_ or ""):
        return None
    p = dossier_evenements(cfg) / f"{id_}.json"
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def reinitialiser() -> None:
    """Vide les caches (tests)."""
    with _verrou:
        _objets_couches.clear()
        _candidats.clear()
        _parts_courbes.clear()
    crise_local.reinitialiser()
    crise_externe.reinitialiser()


# ---------------------------------------------------------------- export Excel (lot C2)

def _nom_feuille(texte: str, pris: set[str]) -> str:
    """Nom de feuille Excel : 31 caractères au plus, sans caractère interdit, unique."""
    base = re.sub(r"[\[\]:*?/\\]", " ", texte).strip()[:31] or "Feuille"
    nom, i = base, 2
    while nom.lower() in pris:
        nom = f"{base[:28]} {i}"
        i += 1
    pris.add(nom.lower())
    return nom


def _heure(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone().strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return iso


def classeur(cfg: dict, evt: dict) -> bytes:
    """Classeur de l'événement : synthèse, une feuille par catégorie (éléments touchés et décisions), main courante.

    Les éléments sont recalculés à partir du point et des zones de l'événement ; les décisions (statuts) viennent
    de l'événement. Un élément décidé qui n'est plus dans le périmètre est gardé, signalé « hors périmètre ».
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not evt.get("centre"):
        raise ErreurCrise("événement sans point : rien à exporter")
    r = analyser(cfg, evt["centre"], evt.get("zones") or [])
    c = cfg["crise"]
    zones = {z["id"]: z for z in c["zones"]}
    libelles_statuts = {x["id"]: x["libelle"] for x in statuts(cfg)}
    defaut = statuts(cfg)[0]["id"]
    decisions = evt.get("statuts") or {}
    gras = Font(bold=True)
    blanc = Font(bold=True, color="FFFFFF")
    entete = PatternFill("solid", fgColor="0A0082")

    def remplir(ws, lignes, largeurs):
        for i, ligne in enumerate(lignes, 1):
            ws.append(ligne)
            if i == 1:
                for cell in ws[1]:
                    cell.font, cell.fill = blanc, entete
        for i, l in enumerate(largeurs, 1):
            ws.column_dimensions[get_column_letter(i)].width = l
        ws.freeze_panes = "A2"
        if len(lignes) > 1:
            ws.auto_filter.ref = ws.dimensions

    def teinter(cell, role):
        z = zones.get(role)
        if z:
            cell.fill = PatternFill("solid", fgColor=z["couleur"].lstrip("#").upper())

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    exercice = evt.get("mode") == "exercice"
    ws.append(["EXERCICE" if exercice else "ÉVÉNEMENT RÉEL"])
    ws["A1"].font = Font(bold=True, size=14, color="1971C2" if exercice else "C92A2A")
    for k, v in [("Événement", evt.get("titre")), ("Type", evt.get("type")), ("Début", _heure(evt.get("debut"))),
                 ("Description", evt.get("description")), ("Auteur", evt.get("auteur")),
                 ("Point (WGS84)", f"{r['centre']['wgs84'][1]:.6f} N, {r['centre']['wgs84'][0]:.6f} E"),
                 ("Point (Lambert 93)", f"X {r['centre']['l93'][0]:.2f}  Y {r['centre']['l93'][1]:.2f}"),
                 ("Export", datetime.now().strftime("%d/%m/%Y %H:%M"))]:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = gras
    ws.append([])
    ws.append(["Zones"])
    ws.cell(ws.max_row, 1).font = gras
    for z in r["zones"]:
        ws.append([zones[z["role"]]["libelle"], f"cercle de {z['rayon_m']:g} m" if z["forme"] == "cercle"
                   else f"couloir de {z['longueur_m']:,} m, {z['tampon_m']:g} m de part et d'autre".replace(",", " ")
                   if z["forme"] == "couloir" else f"polygone de {z['surface_m2'] / 1e4:.2f} ha"])
        teinter(ws.cell(ws.max_row, 1), z["role"])
    ws.append([])
    actifs = [z["id"] for z in c["zones"] if any(x["role"] == z["id"] for x in r["zones"])]
    ws.append(["Catégorie", *[zones[i]["libelle"] for i in actifs], "Décidés"])
    for cell in ws[ws.max_row]:
        cell.font = gras
    for k in r["categories"]:
        if k["role"] == "denombrement":   # estimations (population, bâtiments) : valeurs par zone, pas d'éléments
            if k.get("en_cours"):
                ws.append([k["libelle"], "chargement en cours : relancer l'export dans quelques secondes"])
            for m in k.get("mesures", []):
                a = m.get("arrondi") or 1
                ws.append([f"{k['libelle']} : {m['libelle']}", *[round(m["par_zone"].get(i, 0) / a) * a for i in actifs]])
            continue
        n = [sum(1 for e in k["elements"] if e["zone"] == i) for i in actifs]
        decides = sum(1 for e in k["elements"] if e["zone"]
                      and (decisions.get(f"{k['id']}|{e['cle']}") or {}).get("statut", defaut) != defaut)
        ws.append([k["libelle"], *n, decides])
    ws.append([])
    ws.append(["Avertissement", c.get("avertissement", "")])
    ws.append(["Sources", "; ".join(sorted({f"{s['libelle']} ({_heure(s['date'])})" if s.get("date") else s["libelle"]
                                            for s in r["sources"]}))])
    for a in r["anomalies"]:
        ws.append(["Anomalie", a])
    ws.column_dimensions["A"].width = 30
    for col in "BCDEF":
        ws.column_dimensions[col].width = 22

    pris = {"synthèse", "main courante"}
    for k in r["categories"]:
        lignes = [["Zone", "Élément", "Détail", "Cause", "Distance (m)", "Direction", "Proposition",
                   "Statut", "Décidé le", "Par", "Longitude", "Latitude"]]
        roles_lignes = []
        vus = set()
        for e in k["elements"]:
            d = decisions.get(f"{k['id']}|{e['cle']}") or {}
            if not e["zone"] and not d and k["role"] != "secours":
                continue
            vus.add(e["cle"])
            hors = "hors zone" if k["role"] == "secours" else "hors périmètre"
            lignes.append([zones[e["zone"]]["libelle"] if e["zone"] else hors, e["libelle"], e.get("detail"),
                           e.get("cause"), e["distance_m"], e["direction"], e.get("proposition"),
                           libelles_statuts.get(d.get("statut", defaut), d.get("statut")), _heure(d.get("heure")),
                           d.get("auteur"), e["position"][0], e["position"][1]])
            roles_lignes.append(e["zone"])
        # décisions sur des éléments sortis des candidats (point déplacé au-delà de rayon_max)
        for cle, d in decisions.items():
            cat, _, cle_e = cle.partition("|")
            if cat == k["id"] and cle_e not in vus:
                lignes.append(["hors périmètre", d.get("libelle") or cle_e, None, None, None, None, None,
                               libelles_statuts.get(d.get("statut"), d.get("statut")), _heure(d.get("heure")),
                               d.get("auteur"), None, None])
                roles_lignes.append(None)
        if len(lignes) == 1:
            continue
        f = wb.create_sheet(_nom_feuille(k["libelle"], pris))
        remplir(f, lignes, [16, 40, 30, 34, 12, 10, 40, 18, 16, 14, 11, 11])
        for i, role in enumerate(roles_lignes, 2):
            teinter(f.cell(i, 1), role)
            f.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")

    j = wb.create_sheet("Main courante")
    remplir(j, [["Heure", "Auteur", "Action", "Texte"]]
            + [[_heure(x.get("heure")), x.get("auteur"), x.get("type"), x.get("texte")] for x in evt.get("journal") or []],
            [17, 16, 14, 100])
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()
